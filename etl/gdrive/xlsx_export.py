"""
Вивантаження кодифікатора та реєстру документів у Excel (openpyxl, write_only — тягне 600k+ рядків).
"""
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

_HEAD_FONT = Font(bold=True, color='FFFFFF')
_HEAD_FILL = PatternFill('solid', fgColor='1F4E78')
_WRAP = Alignment(wrap_text=True, vertical='top')

README = [
    ('Кодифікатор документів клієнтів (Google Диск → поля CRM)', ''),
    ('', ''),
    ('ZV-BB', 'Блок справи (сегмент) — 01 … 26 з каталогу полів CRM, 00 — службові / нерозпізнані файли'),
    ('ZV-BB.TT', 'Тип документа всередині блоку, напр. ZV-14.02 = Кредитний звіт УБКІ'),
    ('ZV-F-BB.NNN', 'Код поля CRM: блок поля + порядковий номер поля в блоці'),
    ('ZV-BB.TT-КЛІЄНТ-NNNN', 'Код екземпляра документа в реєстрі: тип + ключ папки клієнта (8 симв.) + № документа '
                             'цього типу в клієнта (за датою створення)'),
    ('', ''),
    ('Зв’язок scan', 'file-поле CRM, куди кладеться скан саме цього документа'),
    ('Зв’язок evidence', 'поле CRM, яке заповнюється / підтверджується цим документом (за колонкою '
                         '«Первинний документ-доказ» каталогу)'),
    ('Зв’язок output', 'для вихідних документів справи (Заява, Декларація, План, Опис майна, Перелік кредиторів, '
                       'Процесуальні) — поля, позначені ✔ у відповідній колонці каталогу'),
    ('', ''),
    ('parse', 'Підказка парсеру: ocr — скан/фото, потрібне розпізнавання; text — текстовий документ (docx/gdoc); '
              'table — таблиця; unzip — архів, тип визначено за назвою, парсити після розпакування; skip — не парсити (службове / чутливе / вихідний запит)'),
    ('sensitive', 'Чутливі дані (ключі ЕЦП, паролі) — вміст не читати й не передавати в CRM'),
    ('confidence', 'Впевненість класифікації: 0.95 назва+папка; 0.9 назва+блок папки; 0.8 тільки назва; '
                   '0.7 назва суперечить папці; 0.65/0.5 тільки папка; 0.5 тільки кредитор у назві; '
                   '0.2 тільки тип файлу'),
]


def _sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[_col(i)].width = w
    ws.freeze_panes = 'A2'
    head = []
    for h in header:
        c = WriteOnlyCell(ws, value=h)
        c.font, c.fill, c.alignment = _HEAD_FONT, _HEAD_FILL, _WRAP
        head.append(c)
    ws.append(head)
    for r in rows:
        ws.append([_cell(v) for v in r])
    return ws


def _col(i):
    s = ''
    i += 1
    while i:
        i, rem = divmod(i - 1, 26)
        s = chr(65 + rem) + s
    return s


def _cell(v):
    if isinstance(v, (list, tuple, set)):
        return '; '.join(str(x) for x in v)
    if isinstance(v, bool):
        return 'так' if v else ''
    if isinstance(v, str):
        # назви файлів з Диска — недовірений текст: «=HYPERLINK(...)» не має ставати формулою
        v = ILLEGAL_CHARACTERS_RE.sub('', v)
        if v.startswith(('=', '+', '-', '@')):
            return "'" + v
    return v


def write_codebook_xlsx(cb: dict, path: str):
    wb = Workbook(write_only=True)
    _add_codebook_sheets(wb, cb)
    wb.save(path)


def _add_codebook_sheets(wb, cb):
    _sheet(wb, 'Як читати', ['Код / поле', 'Значення'], README, [26, 120])

    by_block = Counter(d['block'] for d in cb['doc_types'])
    fields_by_block = Counter(f['block'] for f in cb['fields'])
    _sheet(wb, 'Блоки (сегменти)', ['Код блоку', 'Блок справи', 'Типів документів', 'Полів CRM у блоці'],
           [(b['block_code'], b['name'], by_block.get(b['block'], 0), fields_by_block.get(b['block'], 0))
            for b in cb['blocks']], [12, 40, 18, 18])

    _sheet(wb, 'Типи документів',
           ['Код документа', 'Блок', 'Назва блоку', 'Тип документа', 'Вид', 'Парсинг', 'Чутливе',
            'file-полів (scan)', 'Полів-доказів', 'Полів вихідного док.', 'file-поля CRM',
            'Правила: назва файлу', 'Правила: назва папки'],
           [(d['doc_code'], f"ZV-{d['block']}", d['block_name'], d['name'], d['kind'], d['parse'], d['sensitive'],
             d['n_scan_fields'], d['n_evidence_fields'], d['n_output_fields'], d['scan_fields'],
             d['patterns'], d['folder_patterns']) for d in cb['doc_types']],
           [12, 8, 26, 60, 10, 8, 8, 10, 10, 10, 30, 60, 40])

    _sheet(wb, 'Поля CRM (коди)',
           ['Код поля ZV', 'Код поля CRM', 'Назва поля в CRM', 'Тип', 'Блок', 'Роль поля', 'Що фіксує',
            'Характер даних', 'Первинний документ-доказ', 'Використовується в', 'Макс. % заповнення', 'Сутності CRM'],
           [(f['field_code'], f['crm_code'], f['crm_name'], f['crm_type'], f['block'], f['role'], f['fixes'],
             f['data_nature'], f['evidence'], f['used_in'], f['max_fill_pct'], f['entities'])
            for f in cb['fields']],
           [14, 26, 36, 12, 6, 30, 40, 30, 50, 30, 10, 50])

    names = {d['doc_type']: d['name'] for d in cb['doc_types']}
    _sheet(wb, 'Документ → поля',
           ['Код документа', 'Тип документа', 'Зв’язок', 'Код поля ZV', 'Код поля CRM', 'Назва поля в CRM',
            'Що фіксує', 'Характер даних', 'Макс. % заповнення', 'Сутності CRM'],
           [(l['doc_code'], names[l['doc_type']], l['link_type'], l['field_code'], l['crm_code'], l['crm_name'],
             l['fixes'], l['data_nature'], l['max_fill_pct'], l['entities']) for l in cb['links']],
           [12, 50, 9, 14, 26, 36, 40, 30, 10, 50])

    by_type = defaultdict(lambda: Counter())
    for l in cb['links']:
        by_type[l['doc_type']][l['link_type']] += 1
    _sheet(wb, 'Прогалини CRM',
           ['Код документа', 'Тип документа', 'Звʼязків з полями', 'Первинний документ-доказ без полів у CRM (блок)'],
           [(d['doc_code'], d['name'], sum(by_type[d['doc_type']].values()),
             '; '.join(f"{g['block']}: {g['evidence']}" for g in cb.get('gaps', []) if g['doc_type'] == d['doc_type']))
            for d in cb['doc_types']
            if d['kind'] == 'evidence' and (not by_type[d['doc_type']]
                                            or any(g['doc_type'] == d['doc_type'] for g in cb.get('gaps', [])))],
           [12, 60, 12, 90])


REGISTRY_COLUMNS = [
    ('doc_instance_code', 'Код документа (екземпляр)'),
    ('doc_code', 'Код типу'),
    ('block', 'Блок'),
    ('block_name', 'Сегмент (блок справи)'),
    ('doc_type_name', 'Тип документа'),
    ('confidence', 'Впевненість'),
    ('method', 'Метод'),
    ('rule', 'Правило'),
    ('also_contains', 'Також містить (коди)'),
    ('parse', 'Парсинг'),
    ('sensitive', 'Чутливе'),
    ('subject', 'Чий документ'),
    ('creditor', 'Кредитор'),
    ('is_archived', 'Старе / неактуальне'),
    ('not_for_court', 'Не в суд'),
    ('needs_check', 'Перевірити'),
    ('is_copy', 'Копія / дубль назви'),
    ('is_duplicate', 'Дубль за вмістом (md5)'),
    ('signed_container', 'Підписаний контейнер'),
    ('is_zip', 'Архів (розпакувати)'),
    ('client_key', 'Ключ клієнта'),
    ('client_folder', 'Папка клієнта'),
    ('client_phone', 'Телефон клієнта'),
    ('root_name', 'Коренева папка'),
    ('path', 'Шлях'),
    ('file_name', 'Назва файлу'),
    ('mime_type', 'MIME'),
    ('size_bytes', 'Розмір, байт'),
    ('created_time', 'Створено'),
    ('modified_time', 'Змінено'),
    ('file_id', 'ID файлу Google'),
    ('web_view_link', 'Посилання'),
    ('crm_fields', 'Поля CRM для парсингу (коди ZV)'),
]


def summarize(records):
    """(segments, clients) для аркушів зведення — з записів у пам'яті (офлайн-режим)."""
    seg = defaultdict(lambda: [0, 0, 0.0, set()])
    clients = defaultdict(Counter)
    for r in records:
        s = seg[(r['doc_code'], r['block_name'], r['doc_type_name'])]
        s[0] += 1
        s[1] += 1 if r['confidence'] >= 0.8 else 0
        s[2] += r['confidence']
        s[3].add(r['client_key'])
        clients[(r['client_key'], r['client_folder'], r['client_phone'])][r['block']] += 1
    segments = [(k[0], k[1], k[2], v[0], v[1], len(v[3]), round(v[2] / v[0], 2)) for k, v in sorted(seg.items())]
    client_rows = [(k[0], k[1], k[2], dict(c)) for k, c in sorted(clients.items(), key=lambda x: x[0][1] or '')]
    return segments, client_rows


def write_registry_xlsx(rows, segments, clients, cb: dict, path: str):
    """
    rows     — ітератор dict-ів реєстру (ключі REGISTRY_COLUMNS), може бути потоковим
    segments — [(код, блок, тип, документів, впевнено, клієнтів, сер. впевненість)]
    clients  — [(ключ, папка, телефон, {блок: к-сть})]
    """
    wb = Workbook(write_only=True)
    _sheet(wb, 'Зведення по сегментах',
           ['Код типу', 'Сегмент (блок)', 'Тип документа', 'Документів', 'З них впевнено (≥0.8)', 'Клієнтів',
            'Сер. впевненість'], segments, [12, 30, 60, 12, 14, 10, 12])

    blocks = sorted({b for *_, c in clients for b in c})
    _sheet(wb, 'Клієнти × блоки', ['Ключ клієнта', 'Папка клієнта', 'Телефон', 'Усього'] + [f'ZV-{b}' for b in blocks],
           [(k, folder, phone, sum(c.values())) + tuple(c.get(b, 0) for b in blocks)
            for k, folder, phone, c in clients],
           [12, 40, 16, 8] + [7] * len(blocks))

    _sheet(wb, 'Реєстр документів', [h for _, h in REGISTRY_COLUMNS],
           ([r.get(k) for k, _ in REGISTRY_COLUMNS] for r in rows),
           [26, 10, 6, 26, 45, 8, 14, 24, 14, 7, 7, 8, 18, 7, 7, 7, 7, 7, 7, 10, 30, 14, 14, 50, 40, 20,
            7, 10, 18, 18, 30, 40, 60])

    _add_codebook_sheets(wb, cb)
    wb.save(path)
