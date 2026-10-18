#!/usr/bin/env python3
"""
Каталог полів CRM з юридичним сенсом → Excel.

Джерела полів:
  --audit-xlsx PATH   вивантаження аудиту полів (bitrix_field_audit.xlsx: аркуші Ліди / Продажа /
                      Початок шлях до суду / Суд, колонки «Назва поля», «Код поля», «Тип», «% заповнення»)
  --from-b24          живе вивантаження з Bitrix24: поля + % заповнення по всіх записах
                      (--entities lead,deal,contact,company,sp; потрібні B24_DOMAIN, B24_USER_ID
                      і B24_TOKEN_FIELDS або B24_TOKEN_DEALS)

Приклади:
  python -m tools.crm_fields.build_catalog --audit-xlsx bitrix_field_audit.xlsx
  python -m tools.crm_fields.build_catalog --from-b24 --entities contact,sp --since 2025-06-01 \
      --out docs/crm_fields/crm_fields_contacts_sp.xlsx
"""
import argparse
import re
import os
from collections import OrderedDict, defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .legal_rules import (
    TARGET_DOCS, TOPICS, classify,
    R_EVIDENCE, R_REQUISITES, R_CALC, R_CONTROL, R_SCREEN, R_NOTE, R_PROCESS,
    R_TEMPLATE, R_SERVICE, R_SYSTEM, R_SECRET, R_UNKNOWN,
)

# Аркуш аудиту → коротка назва воронки
AUDIT_SHEETS = OrderedDict([
    ('Ліди', 'Лід'),
    ('Продажа', 'Продаж'),
    ('Початок шлях до суду', 'Підготовка до суду'),
    ('Суд', 'Суд'),
])

ROLE_FILL = {
    R_EVIDENCE: 'C6EFCE', R_REQUISITES: 'E2F0D9', R_CALC: 'DDEBF7', R_PROCESS: 'FFF2CC',
    R_CONTROL: 'EDEDED', R_NOTE: 'F2F2F2', R_SCREEN: 'FCE4D6', R_TEMPLATE: 'F2F2F2',
    R_SERVICE: 'F2F2F2', R_SYSTEM: 'F2F2F2', R_SECRET: 'FF9999', R_UNKNOWN: 'FFFF99',
}
LEGAL_ROLES = {R_EVIDENCE, R_REQUISITES, R_CALC, R_PROCESS}


# --- Завантаження полів ---

def load_from_audit(path: str) -> list:
    """Поля з аудит-файлу: [{code, name, type, entities: {воронка: % заповнення}}]."""
    wb = load_workbook(path, read_only=True)
    fields = OrderedDict()
    for sheet, label in AUDIT_SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        for row in wb[sheet].iter_rows(min_row=4, values_only=True):
            if not row or not row[2] or row[2] == 'Код поля':
                continue
            code = str(row[2]).strip()
            f = fields.setdefault(code, dict(code=code, name=row[1] or code, type=row[3] or '', entities=OrderedDict()))
            pct = row[7]
            f['entities'][label] = float(pct) if isinstance(pct, (int, float)) else None
    return list(fields.values())


def _is_filled(value) -> bool:
    if value is None or value is False:
        return False
    if isinstance(value, str):
        return value.strip() not in ('', '0', 'N', '|UAH', '0|UAH')
    if isinstance(value, (list, dict)):
        return any(_is_filled(v) for v in (value.values() if isinstance(value, dict) else value))
    if isinstance(value, (int, float)):
        return value != 0
    return True


def _fill_rates(items: list, keys: list) -> dict:
    """% заповнення по кожному ключу (None, якщо записів немає)."""
    if not items:
        return {k: None for k in keys}
    return {k: round(100.0 * sum(1 for it in items if _is_filled(it.get(k))) / len(items), 1) for k in keys}


def load_from_b24(entities: set, since: str = None, with_fill: bool = True) -> list:
    """Поля з Bitrix24 REST; з with_fill — ще й % заповнення по всіх записах сутності.

    entities: підмножина {'lead', 'deal', 'contact', 'company', 'sp'}.
    since:    'YYYY-MM-DD' — рахувати заповненість лише по записах, змінених після дати.
    """
    from etl.bitrix.b24 import B24
    token = os.environ.get('B24_TOKEN_FIELDS') or os.environ['B24_TOKEN_DEALS']
    domain, user_id = os.environ['B24_DOMAIN'], int(os.environ['B24_USER_ID'])
    # у змінну іноді кладуть повний URL вебхука: https://<домен>/rest/<user>/<токен>/
    m = re.match(r'https?://([^/]+)/rest/(\d+)/([^/]+)', token.strip())
    if m:
        domain, user_id, token = m.group(1), int(m.group(2)), m.group(3)
    b24 = B24(domain, user_id, token)

    fields = OrderedDict()

    def add(entity, code, meta, pct=None):
        name = meta.get('listLabel') or meta.get('formLabel') or meta.get('title') or code
        f = fields.setdefault(code, dict(code=code, name=name, type=meta.get('type', ''), entities=OrderedDict()))
        f['entities'][entity] = pct

    classic = [('lead', 'crm.lead', 'Лід'), ('deal', 'crm.deal', 'Угода'),
               ('contact', 'crm.contact', 'Контакт'), ('company', 'crm.company', 'Компанія')]
    for key, prefix, label in classic:
        if key not in entities:
            continue
        meta_all = b24.call(f'{prefix}.fields')
        rates = {}
        if with_fill:
            flt = {'>=DATE_MODIFY': f'{since}T00:00:00'} if since else {}
            items = b24.get_list(f'{prefix}.list', b24_filter=flt, select=['*', 'UF_*', 'PHONE', 'EMAIL'])
            rates = _fill_rates(items, list(meta_all))
            label = f'{label} ({len(items)} зап.)'
        for code, meta in meta_all.items():
            add(label, code, meta, rates.get(code))

    if 'sp' in entities:
        for t in b24.call('crm.type.list').get('types', []):
            etid = t['entityTypeId']
            meta_all = b24.call('crm.item.fields', {'entityTypeId': etid}).get('fields', {})
            rates, label = {}, f'СП: {t["title"]}'
            if with_fill:
                flt = {'>=updatedTime': f'{since}T00:00:00'} if since else {}
                items = b24.get_list('crm.item.list', b24_filter=flt, select=['*'], entity_type_id=etid)
                rates = _fill_rates(items, list(meta_all))
                label = f'{label} ({len(items)} зап.)'
            for code, meta in meta_all.items():
                add(label, meta.get('upperName') or code, meta, rates.get(code))
    return list(fields.values())


# --- Побудова каталогу ---

def build_rows(fields: list) -> list:
    by_name = defaultdict(list)
    for f in fields:
        by_name[(f['name'] or '').strip().lower()].append(f['code'])

    rows = []
    for i, f in enumerate(fields, 1):
        c = classify(f['code'], f['name'], f['type'])
        pcts = [p for p in f['entities'].values() if p is not None]
        max_pct = max(pcts) if pcts else None

        recs = []
        if c['role'] == R_SECRET:
            recs.append('Прибрати з CRM; доступи зберігати в менеджері паролів')
        if max_pct is not None and max_pct == 0:
            recs.append('Не заповнюється в жодній воронці — кандидат на архів / видалення')
        elif max_pct is not None and max_pct < 1 and c['role'] in LEGAL_ROLES:
            recs.append('Майже не заповнюється (<1%) — перевірити, чи дані беруться з іншого поля')
        twins = [x for x in by_name[(f['name'] or '').strip().lower()] if x != f['code']]
        if twins:
            recs.append('Дубль назви: ' + ', '.join(twins[:5]) + (' …' if len(twins) > 5 else ''))
        low = (f['name'] or '').lower()
        if ('на видалення' in low or low.startswith(('старе', '[архів', 'видалити', 'тест'))
                or low in ('нове поле', 'новий список')):
            recs.append('Позначене як застаріле')

        rows.append(dict(
            n=i, code=f['code'], name=f['name'], type=f['type'],
            where='; '.join(f'{e}: {p:.1f}%' if p is not None else e for e, p in f['entities'].items()),
            max_pct=max_pct, **c, rec='; '.join(recs),
        ))
    return rows


# --- Excel ---

def _style_header(ws, headers, widths):
    ws.append(headers)
    for col, w in enumerate(widths, 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='305496')
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = 'D2'
    ws.row_dimensions[1].height = 45


def write_xlsx(rows: list, out: str, source_note: str):
    wb = Workbook()

    # 1. Основний реєстр
    ws = wb.active
    ws.title = 'Поля → юридичний сенс'
    headers = ['№', 'Код поля', 'Назва поля в CRM', 'Тип', 'Де є (воронка: % заповнення)', 'Блок справи',
               'Роль поля', 'Що фіксує поле', 'Юридичне значення (що доводить)', 'Характер даних у полі',
               'Первинний документ-доказ',
               *TARGET_DOCS, 'Норма (орієнтир)', 'Рекомендація', 'Примітка']
    widths = [6, 26, 48, 13, 34, 30, 26, 50, 70, 40, 45] + [11] * len(TARGET_DOCS) + [30, 40, 35]
    _style_header(ws, headers, widths)
    wrap = Alignment(wrap_text=True, vertical='top')
    for r in rows:
        ws.append([r['n'], r['code'], r['name'], r['type'], r['where'], r['block'], r['role'], r['what'], r['sense'],
                   r['aspect'], r['evidence'], *['✔' if d in r['docs'] else '' for d in TARGET_DOCS],
                   r['norm'], r['rec'], r['note']])
        row = ws.max_row
        for col in range(1, len(headers) + 1):
            ws.cell(row=row, column=col).alignment = wrap
        ws.cell(row=row, column=7).fill = PatternFill('solid', fgColor=ROLE_FILL.get(r['role'], 'FFFFFF'))
        for col in range(12, 12 + len(TARGET_DOCS)):
            ws.cell(row=row, column=col).alignment = Alignment(horizontal='center', vertical='top')
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{ws.max_row}'

    # 2. Документ → поля (тільки юридично значущі)
    ws2 = wb.create_sheet('Документ → поля')
    _style_header(ws2, ['Документ', 'Блок справи', 'Роль поля', 'Код поля', 'Назва поля в CRM', 'Макс. % заповнення',
                        'Що фіксує поле', 'Первинний документ-доказ'], [22, 30, 26, 26, 50, 12, 60, 45])
    for d in TARGET_DOCS:
        for r in sorted((r for r in rows if d in r['docs'] and r['role'] in LEGAL_ROLES),
                        key=lambda r: (r['block'], -(r['max_pct'] or 0))):
            ws2.append([d, r['block'], r['role'], r['code'], r['name'], r['max_pct'], r['what'], r['evidence']])
            for col in range(1, 9):
                ws2.cell(row=ws2.max_row, column=col).alignment = wrap
    ws2.auto_filter.ref = f'A1:H{ws2.max_row}'

    # 3. Довідник блоків
    ws3 = wb.create_sheet('Довідник блоків')
    _style_header(ws3, ['Блок справи', 'Що підтверджує', 'Первинні документи-докази', 'Куди йде', 'Норма',
                        'Полів у CRM'], [32, 80, 55, 40, 35, 12])
    counts = defaultdict(int)
    for r in rows:
        counts[r['block']] += 1
    seen = set()
    for key, (block, sense, evidence, docs, norm, _) in sorted(TOPICS.items(), key=lambda kv: kv[1][0]):
        if block in seen:
            continue
        seen.add(block)
        related = [t for t in TOPICS.values() if t[0] == block]
        ws3.append([block,
                    '\n'.join('• ' + t[1] for t in related),
                    '\n'.join(sorted({t[2] for t in related if t[2] != '—'})),
                    ', '.join(sorted({d for t in related for d in t[3]}, key=TARGET_DOCS.index)),
                    '; '.join(sorted({t[4] for t in related if t[4] != '—'})),
                    counts.get(block, 0)])
        for col in range(1, 7):
            ws3.cell(row=ws3.max_row, column=col).alignment = wrap

    # 4. Зведення
    ws4 = wb.create_sheet('Зведення')
    ws4.column_dimensions['A'].width = 45
    ws4.column_dimensions['B'].width = 14
    ws4.append(['Каталог полів CRM — юридичний сенс'])
    ws4['A1'].font = Font(bold=True, size=14)
    ws4.append([source_note])
    ws4.append(['Унікальних полів', len(rows)])
    ws4.append(['Записів поле × воронка', sum(len(r['where'].split('; ')) for r in rows)])
    ws4.append([])
    ws4.append(['Роль поля', 'Полів'])
    ws4.cell(row=ws4.max_row, column=1).font = Font(bold=True)
    role_cnt = defaultdict(int)
    for r in rows:
        role_cnt[r['role']] += 1
    for role, n in sorted(role_cnt.items(), key=lambda kv: -kv[1]):
        ws4.append([role, n])
        ws4.cell(row=ws4.max_row, column=1).fill = PatternFill('solid', fgColor=ROLE_FILL.get(role, 'FFFFFF'))
    ws4.append([])
    ws4.append(['Документ', 'Юридично значущих полів'])
    ws4.cell(row=ws4.max_row, column=1).font = Font(bold=True)
    for d in TARGET_DOCS:
        ws4.append([d, sum(1 for r in rows if d in r['docs'] and r['role'] in LEGAL_ROLES)])
    ws4.append([])
    ws4.append(['Рекомендації', 'Полів'])
    ws4.cell(row=ws4.max_row, column=1).font = Font(bold=True)
    for label, pred in [
        ('Чутливі дані (паролі / ЕЦП) — прибрати', lambda r: r['role'] == R_SECRET),
        ('0% заповнення в усіх воронках', lambda r: r['max_pct'] == 0),
        ('Дублі назв', lambda r: 'Дубль назви' in r['rec']),
        ('Не класифіковано', lambda r: r['role'] == R_UNKNOWN),
    ]:
        ws4.append([label, sum(1 for r in rows if pred(r))])
    ws4.append([])
    ws4.append(['Норми — орієнтир; редакцію КУзПБ / ГПК і розміри ПМ / МЗП перевіряти на дату подання.'])
    wb.move_sheet('Зведення', offset=-3)

    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    wb.save(out)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument('--audit-xlsx')
    src.add_argument('--from-b24', action='store_true')
    ap.add_argument('--entities', default='contact,company,sp',
                    help='для --from-b24: lead,deal,contact,company,sp (за замовчуванням — контакти, компанії, смарт-процеси)')
    ap.add_argument('--since', help='для --from-b24: рахувати заповненість лише по записах, змінених з дати YYYY-MM-DD')
    ap.add_argument('--no-fill', action='store_true', help='для --from-b24: лише список полів, без % заповнення')
    ap.add_argument('--out', default='docs/crm_fields/crm_fields_legal_map.xlsx')
    args = ap.parse_args()

    if args.from_b24:
        entities = {e.strip() for e in args.entities.split(',') if e.strip()}
        fields = load_from_b24(entities, since=args.since, with_fill=not args.no_fill)
        note = f'Джерело: Bitrix24 REST ({", ".join(sorted(entities))})' + (f', записи з {args.since}' if args.since else '')
    else:
        fields = load_from_audit(args.audit_xlsx)
        note = 'Джерело: bitrix_field_audit (Ліди / Продаж / Підготовка до суду / Суд, період 2025-06-01 → 2026-02-27)'

    rows = build_rows(fields)
    write_xlsx(rows, args.out, note)
    unknown = sum(1 for r in rows if r['role'] == R_UNKNOWN)
    print(f'{len(rows)} полів → {args.out} (не класифіковано: {unknown})')


if __name__ == '__main__':
    main()
